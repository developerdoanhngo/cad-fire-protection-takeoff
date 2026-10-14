"""Xuất kết quả bóc tách ra Excel, trình bày theo nhận diện DC Software."""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .catalog import SYSTEM_NAME
from .paths import ui_dir

BRAND = "055FD1"
NAVY = "0B1F4A"
LIGHT = "EDF3FF"
ZEBRA = "F6F7FB"
FONT = "Arial"

thin = Side(style="thin", color="D5DCE8")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
H_FILL = PatternFill("solid", fgColor=NAVY)
SYS_FILL = PatternFill("solid", fgColor=BRAND)
BLD_FILL = PatternFill("solid", fgColor=LIGHT)
Z_FILL = PatternFill("solid", fgColor=ZEBRA)
WARN_FILL = PatternFill("solid", fgColor="FFF4E5")

MANUAL_ITEMS = [
    ("Hệ thống chữa cháy", [
        "Máy bơm chữa cháy động cơ điện", "Máy bơm chữa cháy động cơ diesel", "Máy bơm bù áp",
        "Tủ điều khiển máy bơm", "Bình tích áp", "Công tắc áp lực", "Phụ kiện ống (co, tê, côn, măng sông, bích)",
        "Giá treo, ty ren, đai ôm", "Sơn ống", "Cuộn vòi, lăng phun", "Vật tư phụ",
    ]),
    ("Hệ thống báo cháy", ["Tủ trung tâm báo cháy", "Bộ nguồn dự phòng", "Ống luồn dây", "Vật tư phụ"]),
    ("Hệ thống đèn sự cố và thoát nạn", ["Ổ cắm, hộp đấu nối", "Ống luồn dây", "Vật tư phụ"]),
    ("Hệ thống loa thông báo", ["Thiết bị trung tâm điều khiển loa", "Vật tư phụ"]),
    ("Hệ thống hút khói", [
        "Tủ điện động lực và điều khiển quạt", "Cáp điện chống cháy cấp nguồn quạt", "Máy phát điện dự phòng",
        "Phụ kiện ống gió", "Giá đỡ ống gió", "Vật tư phụ",
    ]),
    ("Chi phí chung", ["Vận chuyển, dàn giáo, xe nâng", "Kiểm định", "An toàn lao động, vệ sinh môi trường"]),
]


def _font(bold=False, color="1F2937", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def _header(ws, row, titles, widths):
    for c, (t, w) in enumerate(zip(titles, widths), start=1):
        cell = ws.cell(row=row, column=c, value=t)
        cell.font = _font(True, "FFFFFF")
        cell.fill = H_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[row].height = 30


def _banner(ws, title, project, ncols):
    logo = ui_dir() / "assets" / "logo-dark.png"
    if logo.exists():
        try:
            img = XLImage(str(logo))
            img.width, img.height = 170, 48
            ws.add_image(img, "A1")
        except Exception:
            pass
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22
    last = get_column_letter(ncols)
    ws.merge_cells(f"C1:{last}1")
    ws.merge_cells(f"C2:{last}2")
    ws["C1"] = title
    ws["C1"].font = _font(True, NAVY, 15)
    ws["C1"].alignment = Alignment(horizontal="right", vertical="center")
    ws["C2"] = f"Dự án: {project['name']}   ·   Xuất lúc {datetime.now():%d/%m/%Y %H:%M}"
    ws["C2"].font = _font(False, "475569", 10)
    ws["C2"].alignment = Alignment(horizontal="right", vertical="center")
    return 4


def _num(cell, unit):
    cell.number_format = "#,##0.0" if unit == "m" else "#,##0"
    cell.alignment = Alignment(horizontal="right", vertical="center")


def export(project: dict, result: dict, out_path: str) -> str:
    wb = Workbook()

    # ---------- Tổng hợp ----------
    ws = wb.active
    ws.title = "Tổng hợp"
    widths = [8, 58, 10, 16, 12, 26]
    row = _banner(ws, "BẢNG KHỐI LƯỢNG BÓC TÁCH", project, len(widths))
    _header(ws, row, ["STT", "Hạng mục", "Đơn vị", "Khối lượng", "Cách ra số", "Ghi chú"], widths)
    first_data = row + 1
    row += 1
    for g in result["groups"]:
        for c in range(1, 7):
            ws.cell(row=row, column=c).fill = SYS_FILL
            ws.cell(row=row, column=c).border = BORDER
        ws.cell(row=row, column=1, value=g["roman"]).font = _font(True, "FFFFFF")
        ws.cell(row=row, column=2, value=g["name"].upper()).font = _font(True, "FFFFFF")
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
        row += 1
        for b in g["buildings"]:
            for c in range(1, 7):
                ws.cell(row=row, column=c).fill = BLD_FILL
                ws.cell(row=row, column=c).border = BORDER
            ws.cell(row=row, column=1, value="*").alignment = Alignment(horizontal="center")
            ws.cell(row=row, column=2, value=b["name"]).font = _font(True, NAVY)
            row += 1
            for n, it in enumerate(b["items"], start=1):
                vals = [n, it["material"], it["unit"], it["qty"], it["source"],
                        f"Trang: {', '.join(str(x) for x in sorted(set(it['pages']))[:12])}"]
                for c, v in enumerate(vals, start=1):
                    cell = ws.cell(row=row, column=c, value=v)
                    cell.border = BORDER
                    cell.font = _font(color="475569" if c in (5, 6) else "1F2937", size=9 if c == 6 else 10)
                    if n % 2 == 0:
                        cell.fill = Z_FILL
                ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
                ws.cell(row=row, column=3).alignment = Alignment(horizontal="center")
                _num(ws.cell(row=row, column=4), it["unit"])
                row += 1
    ws.freeze_panes = ws.cell(row=first_data, column=1)
    row += 1
    ws.cell(row=row, column=2, value="Ghi chú: “Đếm” = phần mềm đếm ký hiệu; “Đo” = phần mềm đo chiều dài trên bản vẽ "
                                     "(đã nhân hệ số tầng, hao hụt theo cài đặt). Hạng mục không có trên bản vẽ xem sheet “Nhập tay”.").font = _font(italic=True, color="64748B", size=9)
    ws.sheet_view.showGridLines = False

    # ---------- Chi tiết theo trang ----------
    wd = wb.create_sheet("Chi tiết theo trang")
    widths = [16, 26, 8, 52, 18, 14, 9, 40, 9, 14, 14]
    row = _banner(wd, "CHI TIẾT KHỐI LƯỢNG THEO TRANG BẢN VẼ", project, len(widths))
    _header(wd, row, ["Hệ", "File", "Trang", "Tiêu đề trang", "Tòa nhà", "Tầng", "Hệ số tầng", "Hạng mục",
                      "Đơn vị", "Trên bản vẽ", "Khối lượng"], widths)
    first = row + 1
    row += 1
    for n, d in enumerate(result["detail"]):
        vals = [SYSTEM_NAME.get(d["system"], d["system"]).replace("Hệ thống ", ""), d["file"], d["page"], d["title"],
                d["building"], d["floor"], d["multiplier"], d["material"], d["unit"], d["base"], d["qty"]]
        for c, v in enumerate(vals, start=1):
            cell = wd.cell(row=row, column=c, value=v)
            cell.border = BORDER
            cell.font = _font(size=9)
            if n % 2:
                cell.fill = Z_FILL
        _num(wd.cell(row=row, column=10), d["unit"])
        _num(wd.cell(row=row, column=11), d["unit"])
        row += 1
    wd.freeze_panes = wd.cell(row=first, column=1)
    wd.auto_filter.ref = f"A{first - 1}:K{max(row - 1, first)}"
    wd.sheet_view.showGridLines = False

    # ---------- Cần kiểm tra ----------
    wk = wb.create_sheet("Cần kiểm tra")
    widths = [6, 70, 80]
    row = _banner(wk, "NHỮNG CHỖ CẦN NGƯỜI KIỂM TRA", project, len(widths))
    _header(wk, row, ["#", "Vị trí", "Nội dung"], widths)
    row += 1
    for n, w in enumerate(result["warnings"], start=1):
        for c, v in enumerate([n, w["where"], w["text"]], start=1):
            cell = wk.cell(row=row, column=c, value=v)
            cell.border = BORDER
            cell.font = _font(size=9)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = WARN_FILL
        row += 1
    excluded = [p for p in project["pages"] if not p["include"]]
    row += 1
    wk.cell(row=row, column=2, value=f"Các trang không tính ({len(excluded)} trang)").font = _font(True, NAVY)
    row += 1
    for p in excluded:
        wk.cell(row=row, column=2, value=f"{p['file_name']} · Trang {p['index']}: {p['title'][:90]}").font = _font(size=9)
        wk.cell(row=row, column=3, value=p.get("auto_reason") or "Người dùng bỏ chọn").font = _font(size=9, color="64748B")
        row += 1
    wk.sheet_view.showGridLines = False

    # ---------- Nhập tay ----------
    wm = wb.create_sheet("Nhập tay")
    widths = [8, 60, 10, 16, 30]
    row = _banner(wm, "HẠNG MỤC NHẬP TAY (KHÔNG CÓ TRÊN BẢN VẼ)", project, len(widths))
    _header(wm, row, ["STT", "Hạng mục", "Đơn vị", "Khối lượng", "Ghi chú"], widths)
    row += 1
    for gname, items in MANUAL_ITEMS:
        for c in range(1, 6):
            wm.cell(row=row, column=c).fill = BLD_FILL
            wm.cell(row=row, column=c).border = BORDER
        wm.cell(row=row, column=2, value=gname).font = _font(True, NAVY)
        row += 1
        for n, it in enumerate(items, start=1):
            for c, v in enumerate([n, it, "", None, ""], start=1):
                cell = wm.cell(row=row, column=c, value=v)
                cell.border = BORDER
                cell.font = _font()
            wm.cell(row=row, column=4).fill = PatternFill("solid", fgColor="FFFBEB")
            row += 1
    wm.sheet_view.showGridLines = False

    # ---------- Ký hiệu ----------
    wsy = wb.create_sheet("Bảng ký hiệu")
    widths = [44, 14, 46, 10, 14]
    row = _banner(wsy, "BẢNG KÝ HIỆU → VẬT TƯ ĐÃ DÙNG", project, len(widths))
    _header(wsy, row, ["Tên ký hiệu trong bản vẽ", "Chữ hiện ra", "Vật tư", "Đơn vị", "Trạng thái"], widths)
    row += 1
    for key, s in sorted(project["symbols"].items()):
        m = s.get("map") or {}
        status = "Bỏ qua" if m.get("ignore") else ("Chưa xác định" if not m.get("material") else
                                                 {"user": "Đã xác nhận", "library": "Từ thư viện", "suggest": "Gợi ý"}.get(s.get("source"), ""))
        for c, v in enumerate([s["eff"], s["label"], m.get("material", ""), m.get("unit", ""), status], start=1):
            cell = wsy.cell(row=row, column=c, value=v)
            cell.border = BORDER
            cell.font = _font(size=9)
        row += 1
    wsy.sheet_view.showGridLines = False

    out = Path(out_path)
    if out.suffix.lower() != ".xlsx":
        out = out.with_suffix(".xlsx")
    wb.save(out)
    return str(out)
